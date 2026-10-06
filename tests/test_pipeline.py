# -*- coding: utf-8 -*-
"""
Pipeline regression tests: validator, specification parsers, the markdown
format, the drafter, and the two build gates.

Several tests are named after a real defect found while building the tool;
they exist so that defect can never return unnoticed.
"""

import copy
import json
import os
import re
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.join(HERE, "..", "engine")
sys.path[:0] = [HERE, ENGINE]
import fixtures as F                                                  # noqa: E402
import probe_layout as P                                              # noqa: E402
from validate_config import validate                                  # noqa: E402
from spec_parse import (parse_game_spec, parse_data_spec, parse_restrictions,  # noqa: E402
                        read_win_rule, forbidden_from_text)
from spec_md import render, parse, unticked_review_items               # noqa: E402
from doc_text import flatten, to_text                                  # noqa: E402
import draft_spec                                                      # noqa: E402


def facts_for(cfg, tickets):
    path = os.path.join(tempfile.mkdtemp(), "layout.TXT")
    F.write_layout(path, cfg, tickets)
    st = P.sniff_structure(path)
    return P.to_facts(st, P.measure(path, st, P.find_void_serial(path, st)))


def normalise(cfg):
    c = copy.deepcopy(cfg)
    c.pop("schema_version", None)
    c["job"].pop("sources", None)
    for g in c["games"]:
        if "channel" in g:                      # older settings: one channel
            g["channels"] = [g.pop("channel")]
        g["prizes"] = sorted(g["prizes"], key=lambda p: p["amount"])
        g["special_symbols"] = g.get("special_symbols") or []
    c["agent"]["low_letters"] = {str(k): v for k, v in c["agent"]["low_letters"].items()}
    return c


# ====================================================================== #
class Validator(unittest.TestCase):

    def setUp(self):
        self.cfg = F.fixture_config()

    def errors(self, cfg):
        return " | ".join(m for _, m in validate(cfg).errors)

    def test_clean_settings_are_valid(self):
        self.assertTrue(validate(self.cfg).ok)

    def test_impossible_batch_rule_is_caught(self):
        """The 0345 defect: a batch promise the prize count cannot meet."""
        self.cfg["games"][0]["group_rules"].append(
            {"group": "batch", "size": 3, "amounts": [1000], "mode": "min", "n": 1})
        self.assertIn("IMPOSSIBLE", self.errors(self.cfg))

    def test_mistyped_prize_count_is_caught_several_ways(self):
        self.cfg["games"][0]["prizes"][2]["count"] = 40
        e = self.errors(self.cfg)
        for phrase in ("expected 400", "prize pool", "payout"):
            self.assertIn(phrase, e)

    def test_packing_not_covering_every_book(self):
        self.cfg["games"][0]["packing"][0]["slots"] = 3
        self.assertIn("packing schedule covers", self.errors(self.cfg))

    def test_gleps_disagreeing_with_prize_table(self):
        self.cfg["games"][0]["gleps_per_book"] = 5
        self.assertIn("GLEPS", self.errors(self.cfg))

    def test_unknown_win_rule(self):
        self.cfg["games"][0]["win_rule"] = "match_three"
        self.assertIn("unknown win rule", self.errors(self.cfg))

    def test_special_symbol_missing_from_table(self):
        self.cfg["games"][1]["special_symbols"] = ["KB", "ZZ"]
        self.assertIn("not in the play-symbol table", self.errors(self.cfg))

    def test_payout_disagreeing_with_declared(self):
        self.cfg["games"][0]["expected"]["payout_pct"] = 53.0
        self.assertIn("payout", self.errors(self.cfg))

    # ---- channels: a game may be printed on several (NLB 0115: four) ----
    def test_one_game_on_several_channels_is_valid(self):
        self.assertTrue(validate(F.fixture_config(split=True)).ok)

    def test_channel_belonging_to_no_game(self):
        self.cfg["layout"]["channels"] = 3
        self.assertIn("channel(s) 3 belong to no game", self.errors(self.cfg))

    def test_channel_claimed_by_two_games(self):
        cfg = F.fixture_config(split=True)
        cfg["games"][1]["channels"] = [2]
        e = self.errors(cfg)
        self.assertIn("channel 2 is assigned to both", e)
        self.assertIn("channel(s) 3 belong to no game", e)

    def test_older_single_channel_settings_are_valid(self):
        for g in self.cfg["games"]:
            g["channel"] = g.pop("channels")[0]
        self.assertTrue(validate(self.cfg).ok)


# ====================================================================== #
class Parsers(unittest.TestCase):

    def test_restriction_wordings(self):
        text = ("A slot of 12,000 Tickets (300 books) should contain the following prizes. "
                "There will be 41 boxes with this prize allocation. • Rs. 500 - at least 300 "
                "winners • Rs. 10,000 - 02 winner(s) • Umbrella (Rs. 1,200) - 75 winners "
                "The last box containing 8,000 Tickets (200 books) should have: "
                "• Rs. 500 - at least 200 winners "
                "Each book should contain one prize of Rs. 500. "
                "Each book should contain either one prize of Rs.200, Rs.500 or both. "
                "Each book should contain either one prize of Rs. 200 or Rs. 500. "
                "Each book should not contain more than one prize of Rs.5,000, Rs.10,000 or Rs.25,000. "
                "Each batch of 75 books should contain a minimum of one Rs. 10,000 winner. "
                "Each book should contain a lucky rabbit. "
                'Each book should contain at least one "NON-GLEPS" prize. '
                "No consecutive tickets with same play patterns. No consecutive tickets with "
                "same winning symbols in any rotation. No logical relationship between play "
                "patterns and VIRN. No more than three consecutive tickets with same prize "
                "value. Any multiples of 5 tickets starting from ticket number 001 in a book "
                "should contain at least one winner.")
        r = parse_restrictions(text, 500000, 40)
        self.assertEqual(r["packing"][0], {"slots": 41, "books": 300, "allocation": [
            {"amount": 500, "mode": "min", "n": 300},
            {"amount": 10000, "mode": "exact", "n": 2},        # "02 winner(s)" is exact
            {"amount": 1200, "mode": "exact", "n": 75}]})      # a labelled prize
        self.assertEqual(r["packing"][1]["slots"], 1)
        modes = [(g["mode"], g["amounts"], g.get("size")) for g in r["group_rules"]]
        self.assertEqual(modes, [
            ("exactly", [500], 1),
            ("at_least_one_total", [200, 500], 1),               # "or both"
            ("exactly_one_total", [200, 500], 1),                # "either ... or"
            ("at_most_each", [5000, 10000, 25000], 1),
            ("min", [10000], 75)])                               # batch
        self.assertEqual(r["unrecognised"], ["Each book should contain a lucky rabbit"])
        self.assertEqual(r["red_missing"], [])
        self.assertEqual(r["fixed"]["max_same_value_run"], 3)
        self.assertEqual(r["fixed"]["winner_every_n_tickets"], 5)

    def test_missing_standard_restriction_is_reported(self):
        r = parse_restrictions("No consecutive tickets with same play patterns.", 400, 20)
        self.assertIn("no play pattern / VIRN relationship", r["red_missing"])

    def test_win_rule_readings(self):
        cases = {
            'Scratch and find two similar symbols or a "Flag" symbol and win':
                ("match_two_or_special", ["Flag"]),
            'Scratch and find a "Gold Coin" symbol or a "Diamond" symbol and win':
                ("special_only", ["Gold Coin", "Diamond"]),
            "Scratch and match two similar money amounts or two Umbrella symbols and win":
                ("match_two", []),                              # a pair, not a special
            "Scratch and find two similar Money amounts and win that corresponding amount":
                ("match_two", []),
            "Scratch and match two similar Numbers and win the corresponding price":
                ("match_two", []),                              # 0112's typo
        }
        for text, want in cases.items():
            self.assertEqual(read_win_rule(text)[:2], want, text)

    def test_forbidden_readings(self):
        self.assertEqual(forbidden_from_text('Two "Flag" symbols'), "two_special")
        self.assertEqual(forbidden_from_text('"Gold Coin" and "Diamond" symbols together'),
                         "two_special")
        self.assertIsNone(forbidden_from_text("N/A"))

    def test_game_spec_ignores_the_contents_page(self):
        """Defect: the mechanism and symbol names were read from the table of
        contents, which carries the same headings."""
        cfg = F.fixture_config()
        g = parse_game_spec(F.game_spec_text(cfg, "0902"), "x")
        self.assertEqual(g["mechanism"], F.MECH["0902"])
        self.assertEqual(g["play_symbol_names"],
                         list(cfg["games"][1]["play_symbols"].values()))
        self.assertEqual(g["problems"], [])

    def test_game_spec_all_fields(self):
        cfg = F.fixture_config()
        g = parse_game_spec(F.game_spec_text(cfg, "0901"), "x")
        self.assertEqual((g["key"], g["name"], g["lottery"], g["ticket_price"]),
                         ("0901", "TEST MATCH", "300", 50))
        self.assertEqual((g["book_size"], g["gleps_amount"], g["gleps_per_book"]), (20, 50, 4))
        self.assertEqual(g["agent_low"], {50: "T"})
        self.assertEqual([(p["sys"], p["amount"], p["count"]) for p in g["prizes"]],
                         [("03", 50, 80), ("07", 500, 20), ("08", 1000, 4)])
        self.assertEqual((g["loss_code"], g["loss_count"]), ("99", 296))
        self.assertEqual((g["doc_version"], g["doc_status"]), ("1.0", "Final"))

    def test_data_spec(self):
        cfg = F.fixture_config()
        d = parse_data_spec(F.data_spec_text(cfg), "x")
        self.assertEqual(d["channel_map"], {"0901": [1], "0902": [2]})
        self.assertEqual(d["agent_low"], {50: "T"})
        self.assertEqual(d["virn_range"], [2000000, 6999999])
        self.assertEqual(d["leading_void_pages"], 2)
        self.assertEqual(d["games"]["0902"]["not_allowed"],
                         '"Gold Coin" and "Diamond" symbols together')

    def test_word_nested_table_is_read(self):
        """Defect: Word specs keep the low-tier agent letter in a table nested
        inside a table cell, which was silently dropped."""
        from docx import Document
        doc = Document()
        outer = doc.add_table(rows=1, cols=2)
        outer.cell(0, 0).text = "Low-Tier Winning Tickets"
        cell = outer.cell(0, 1)
        cell.text = "Two distinct letters chosen from: W, S, A, D, Z, R plus one additional letter:"
        inner = cell.add_table(rows=2, cols=2)
        inner.cell(0, 0).text, inner.cell(0, 1).text = "Prize", "Agent Code Third Letter"
        inner.cell(1, 0).text, inner.cell(1, 1).text = "Rs. 40", "M"
        path = os.path.join(tempfile.mkdtemp(), "nested.docx")
        doc.save(path)
        self.assertIn("Third Letter Rs. 40 M", flatten(to_text(path)))


# ====================================================================== #
class MarkdownFormat(unittest.TestCase):

    def test_round_trip_is_exact(self):
        cfg = F.fixture_config()
        cfg["games"][0]["group_rules"].append(
            {"group": "book", "size": 1, "amounts": [5000, 10000, 25000],
             "mode": "at_most_each", "n": 1})
        self.assertEqual(normalise(parse(render(cfg, [], {}, {"generated": "t"}))),
                         normalise(cfg))

    def test_thousands_separators_survive(self):
        """Defect: '5,000, 10,000' was split at the thousands commas."""
        cfg = F.fixture_config()
        cfg["games"][0]["group_rules"] = [{"group": "book", "size": 1,
                                           "amounts": [5000, 10000, 25000],
                                           "mode": "at_most_each", "n": 1}]
        back = parse(render(cfg, [], {}, {"generated": "t"}))
        self.assertEqual(back["games"][0]["group_rules"][0]["amounts"], [5000, 10000, 25000])

    def test_real_job_settings_round_trip(self):
        jobs = os.path.join(HERE, "..", "jobs")
        found = 0
        for name in sorted(os.listdir(jobs)) if os.path.isdir(jobs) else []:
            p = os.path.join(jobs, name, "job_config.json")
            if os.path.exists(p):
                cfg = json.load(open(p))
                self.assertEqual(normalise(parse(render(cfg, [], {}, {"generated": "t"}))),
                                 normalise(cfg), name)
                found += 1
        if not found:
            self.skipTest("no real job settings present")

    def test_game_on_several_channels_round_trips(self):
        cfg = F.fixture_config(split=True)
        md = render(cfg, [], {}, {"generated": "t"})
        self.assertIn("| 1 / 2 | 0901 |", md)
        self.assertEqual(normalise(parse(md)), normalise(cfg))

    def test_unticked_items_detected(self):
        md = "## Review items\n\n- [x] **R1** done\n- [ ] **R2** open\n\n## 1. Usage\n"
        self.assertEqual(len(unticked_review_items(md)), 1)


# ====================================================================== #
class Drafter(unittest.TestCase):
    """The drafter end to end on the synthetic job: specification text +
    measured facts -> settings, with cross-checks turned into review items."""

    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.tix = F.make_job(seed=7)
        cls.facts = facts_for(cls.cfg, cls.tix)

    def draft(self, docs=None, facts=None):
        docs = docs or self.docs()
        return draft_spec.build("TEST_0901-0902", docs,
                                self.facts if facts is None else facts)

    def docs(self, heading_override=None, extra=None):
        extra = extra or {}
        return [("0901.pdf", F.game_spec_text(self.cfg, "0901", extra.get("0901", ""))),
                ("0902.pdf", F.game_spec_text(self.cfg, "0902", extra.get("0902", ""))),
                ("data.docx", F.data_spec_text(self.cfg, heading_override))]

    def texts(self, D):
        return " | ".join(r["text"] for r in D.review)

    def test_consistent_sources_give_no_review_items(self):
        cfg, D = self.draft()
        self.assertEqual(D.review, [])

    def test_drafted_settings_equal_the_true_settings(self):
        cfg, D = self.draft()
        self.assertEqual(normalise(cfg), normalise(self.cfg))

    def test_full_chain_markdown_to_clean_audit(self):
        from audit_engine import Job, load_layout, run_all_checks
        cfg, D = self.draft()
        built = parse(render(cfg, D.review, D.ev, {"generated": "t"}))
        self.assertTrue(validate(built).ok)
        path = os.path.join(tempfile.mkdtemp(), "layout.TXT")
        F.write_layout(path, built, self.tix)
        job = Job(built)
        games, info = load_layout(job, path)
        for k, g in games.items():
            self.assertEqual([r.index for r in run_all_checks(job, g, info) if not r.passed], [], k)

    def test_mislabelled_symbol_table_is_matched_by_codes(self):
        cfg, D = self.draft(docs=self.docs(heading_override={"0902": "320 0109"}))
        self.assertIn("'Game 320 0109'", self.texts(D))
        self.assertIn("used for 0902", self.texts(D))
        self.assertEqual(cfg["games"][1]["play_symbols"], self.cfg["games"][1]["play_symbols"])

    def test_spec_count_disagreeing_with_data_is_layout_finding(self):
        facts = copy.deepcopy(self.facts)
        facts["games"]["0901"]["system_prize_codes"]["08"] = 5
        cfg, D = self.draft(facts=facts)
        hit = [r for r in D.review if "prize code 08" in r["text"]]
        self.assertTrue(hit and hit[0]["severity"] == "layout")

    def test_win_rule_disagreeing_with_data_is_layout_finding(self):
        facts = copy.deepcopy(self.facts)
        facts["games"]["0902"]["inferred_win_rule"]["inferred_rule"] = "match_two"
        cfg, D = self.draft(facts=facts)
        hit = [r for r in D.review if "behaves as" in r["text"]]
        self.assertTrue(hit and hit[0]["severity"] == "layout")

    def test_impossible_restriction_is_blocking(self):
        cfg, D = self.draft(docs=self.docs(extra={
            "0901": "Each batch of 3 books should contain a minimum of one Rs. 1,000 winner. "}))
        hit = [r for r in D.review if "IMPOSSIBLE" in r["text"]]
        self.assertTrue(hit and hit[0]["severity"] == "blocking")

    def test_unrecognised_restriction_is_blocking(self):
        cfg, D = self.draft(docs=self.docs(extra={
            "0902": "Each book should contain a lucky rabbit. "}))
        self.assertIn("lucky rabbit", self.texts(D))

    def test_no_measurements_is_blocking(self):
        cfg, D = draft_spec.build("T", self.docs(), None)
        self.assertIn("no measured_facts.json", self.texts(D))


class DrafterOnSplitChannels(Drafter):
    """Every drafter test again, with game 0901 printed on two channels.
    Defect (NLB 0115, one game on four channels): the drafter kept only the
    last channel, compared the specification's totals with a quarter of the
    tickets, and raised a false review item for every prize code."""

    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.tix = F.make_job(seed=7, split=True)
        cls.facts = facts_for(cls.cfg, cls.tix)

    def test_split_game_is_noted_with_its_serial_ranges(self):
        cfg, D = self.draft()
        self.assertEqual(cfg["games"][0]["channels"], [1, 2])
        self.assertIn("game 0901 is printed on channels 1, 2 (channel 1: 1 - 200; "
                      "channel 2: 201 - 400)", " ".join(D.ev["layout"]))

    def test_older_facts_without_game_totals_are_blocking(self):
        facts = copy.deepcopy(self.facts)
        del facts["games"]
        cfg, D = self.draft(facts=facts)
        hit = [r for r in D.review if "Re-run probe_layout.py" in r["text"]]
        self.assertTrue(hit and hit[0]["severity"] == "blocking")


# ====================================================================== #
class BuildGates(unittest.TestCase):
    """build_config.py refuses unreviewed or invalid specifications."""

    def run_build(self, md):
        import subprocess
        job = tempfile.mkdtemp()
        open(os.path.join(job, "spec.md"), "w", encoding="utf-8").write(md)
        r = subprocess.run([sys.executable, os.path.join(ENGINE, "build_config.py"),
                            "--job", job], capture_output=True, text=True)
        return r.returncode, r.stdout + r.stderr, job

    def spec(self, review, cfg=None):
        cfg = cfg or F.fixture_config()
        return render(cfg, review, {}, {"generated": "t"})

    def test_unticked_item_blocks_the_build(self):
        rc, out, _ = self.run_build(self.spec([{"id": "R1", "game": None,
                                                "severity": "check", "text": "x"}]))
        self.assertEqual(rc, 1)
        self.assertIn("still unticked", out)

    def test_ticking_without_fixing_is_still_refused(self):
        cfg = F.fixture_config()
        cfg["games"][0]["group_rules"].append(
            {"group": "batch", "size": 3, "amounts": [1000], "mode": "min", "n": 1})
        md = self.spec([{"id": "R1", "game": "0901", "severity": "blocking",
                         "text": "IMPOSSIBLE"}], cfg).replace("- [ ]", "- [x]")
        rc, out, _ = self.run_build(md)
        self.assertEqual(rc, 1)
        self.assertIn("fail validation", out)

    def test_reviewed_valid_spec_builds(self):
        rc, out, job = self.run_build(self.spec([]))
        self.assertEqual(rc, 0, out)
        built = json.load(open(os.path.join(job, "job_config.json")))
        self.assertEqual(normalise(built), normalise(F.fixture_config()))


if __name__ == "__main__":
    unittest.main(verbosity=2)
