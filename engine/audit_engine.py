#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NLB SCRATCH LOTTERY - LAYOUT AUDIT ENGINE  (config-driven)
==========================================================

One tool for every NLB job.  All job-specific numbers live in a settings file
(job_config.json); this file holds only the checking logic and never needs
editing for a new job.

    python audit_engine.py --config <job_config.json> --layout <layout file>
                           [--out <folder>]

WHAT IS FIXED IN CODE (the "red" restrictions - identical in every NLB spec)
---------------------------------------------------------------------------
  * each book contains at least one NON-GLEPS prize
  * no consecutive tickets with the same play pattern
  * no consecutive tickets with the same winning symbols in any rotation
  * no logical relationship between play patterns and VIRN
  * no more than three consecutive tickets with the same prize value
  * any multiple of 5 tickets from ticket 001 in a book contains a winner
The two embedded numbers (3 and 5) are still settings, with those defaults.

WHAT COMES FROM THE SETTINGS (the "yellow" restrictions)
-------------------------------------------------------
  * games, channels, symbols, prize tables
  * book size, packing schedule, per-group prize promises
  * the winning rule and any special symbols
  * GLEPS promise and the agent letter for the low-tier prize

WINNING RULES SUPPORTED
-----------------------
  match_two             two identical play symbols win
  match_two_or_special  two identical, OR a single special symbol
  special_only          a single special symbol wins; a matching pair does NOT

LAYOUT FILE
-----------
Header row, then one row per print page position:
    PageNo + N ticket blocks of 21 fields (one block per channel).  Each
    channel belongs to one game; a game may own several channels, each
    carrying part of its serials.  Tickets are filed by serial number, so
    every check reads a game in serial order whatever its channels.
VOID filler rows carry the configured void serial and are skipped wherever
they appear - leading or trailing padding is normal and is not an exception.
"""

import argparse
import datetime
import json
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from validate_config import validate, game_channels   # noqa: E402


# ------------------------------------------------------------------ #
#  Job wrapper - parsed settings, with derived values computed once   #
# ------------------------------------------------------------------ #

class Job:
    def __init__(self, cfg):
        self.cfg = cfg
        s = cfg["structure"]
        self.tickets = s["tickets_per_game"]
        self.book_size = s["book_size"]
        self.books = self.tickets // self.book_size
        self.price = cfg["job"]["ticket_price"]
        L = cfg["layout"]
        self.block_len = L["block_len"]
        self.channels = L["channels"]
        self.lead = L["lead_columns"]
        self.off = L["offsets"]
        self.void = L["void_serial"]
        self.flag = L["passive_flag"]
        self.virn_lo, self.virn_hi = L["virn_random_range"]
        self.fields = self.lead + self.channels * self.block_len
        A = cfg["agent"]
        self.high_letters = set(A["high_letters"])
        self.low_letters = {int(k): v for k, v in A["low_letters"].items()}
        self.n_agent = A["letters_used"]
        F = cfg["fixed_rules"]
        self.max_run = F["max_same_value_run"]
        self.win_every = F["winner_every_n_tickets"]
        self.gleps_amount = F["gleps_amount"]
        self.games = {g["key"]: g for g in cfg["games"]}
        self.by_channel = {ch: g["key"] for g in cfg["games"]
                           for ch in game_channels(g)}

    def money(self, amount):
        return f"Rs.{amount:,}"


# ------------------------------------------------------------------ #
#  Per-game data, indexed by (serial number - 1)                      #
# ------------------------------------------------------------------ #

class GameData:
    def __init__(self, job, gcfg):
        self.job = job
        self.cfg = gcfg
        self.key = gcfg["key"]
        n = job.tickets
        self.serial = [None] * n
        self.virn = [None] * n
        self.bar = [None] * n
        self.sys = [None] * n
        self.pcode = [None] * n
        self.p1 = [None] * n
        self.p2 = [None] * n
        self.blanks_ok = [True] * n
        self.agent = [None] * n
        # lookups
        self.prize_by_sys = {p["sys"]: p for p in gcfg["prizes"]}
        self.loss_code = gcfg["loss"]["code"]
        self.specials = set(gcfg.get("special_symbols") or [])
        self.symbols = set(gcfg["play_symbols"])

    def present(self, i):
        return self.serial[i] is not None

    def is_winner(self, i):
        return self.sys[i] is not None and self.sys[i] != self.loss_code

    def amount(self, i):
        p = self.prize_by_sys.get(self.sys[i])
        return p["amount"] if p else 0


# ------------------------------------------------------------------ #
#  Loading                                                            #
# ------------------------------------------------------------------ #

def load_layout(job, path):
    games = {k: GameData(job, g) for k, g in job.games.items()}
    info = {
        "rows": 0, "void_tickets": 0, "void_rows_leading": 0,
        "void_rows_trailing": 0, "pages": set(), "bad_field_count": [],
        "malformed": [], "duplicates": [], "wrong_game": [], "qa": {},
    }
    O = job.off
    seen_live = False
    trailing = 0

    with open(path, "r", newline="") as f:
        f.readline()                                   # header
        row_no = 1
        for line in f:
            row_no += 1
            row = line.rstrip("\r\n").split(",")
            info["rows"] += 1
            if len(row) != job.fields:
                if len(info["bad_field_count"]) < 200:
                    info["bad_field_count"].append((row_no, len(row)))
                continue
            if job.lead:
                info["pages"].add(row[0])
            live_here = False
            for ch in range(1, job.channels + 1):
                base = job.lead + (ch - 1) * job.block_len
                serial = row[base + O["serial"]]
                if serial == job.void:
                    info["void_tickets"] += 1
                    continue
                live_here = True
                key = job.by_channel[ch]
                g = games[key]
                lot = g.cfg["lottery"]
                if not (len(serial) == 15 and serial.isdigit()
                        and serial[0] == job.flag and serial[1:4] == lot):
                    info["malformed"].append((row_no, ch, serial)); continue
                if serial[4:8] != key:
                    info["wrong_game"].append((row_no, ch, serial)); continue
                tno = int(serial[8:])
                if not (1 <= tno <= job.tickets):
                    info["malformed"].append((row_no, ch, serial)); continue
                i = tno - 1
                if g.serial[i] is not None:
                    info["duplicates"].append((key, serial)); continue
                g.serial[i] = serial
                g.virn[i] = row[base + O["virn"]]
                g.bar[i] = row[base + O["bar"]]
                g.sys[i] = row[base + O["sys"]]
                g.pcode[i] = row[base + O["print"]]
                pf = base + O["play_first"]
                used = O["play_used"]
                g.p1[i] = row[pf].strip()
                g.p2[i] = row[pf + 1].strip()
                g.blanks_ok[i] = all(
                    row[pf + k].strip() == ""
                    for k in range(used, O["play_count"]))
                af = base + O["agent_first"]
                g.agent[i] = tuple(row[af + k].strip()
                                   for k in range(O["agent_count"]))
                q = row[base + O["qa"]].strip() or "(blank)"
                info["qa"][q] = info["qa"].get(q, 0) + 1
            if live_here:
                seen_live = True
                trailing = 0
            elif seen_live:
                trailing += 1
                info["void_rows_trailing"] = trailing
            else:
                info["void_rows_leading"] += 1
    return games, info


# ------------------------------------------------------------------ #
#  Result object                                                      #
# ------------------------------------------------------------------ #

class CheckResult:
    def __init__(self, index, title, procedure):
        self.index = index
        self.title = title
        self.procedure = procedure
        self.passed = True
        self.finding = []
        self.offenders = []

    def fail(self, items):
        self.passed = False
        self.offenders.extend(str(x) for x in items)

    def note(self, line):
        self.finding.append(line)


def iter_groups(job, books_per_group):
    """Yield (index, start_ticket_idx, end_ticket_idx) for each group of
    `books_per_group` books.  A trailing partial group is included."""
    per = books_per_group * job.book_size
    n = job.tickets
    i = 0
    idx = 0
    while i < n:
        idx += 1
        yield idx, i, min(i + per, n)
        i += per


def count_amount(g, s, e, amount):
    return sum(1 for i in range(s, e) if g.amount(i) == amount)


# ------------------------------------------------------------------ #
#  Winning-rule engine                                                #
# ------------------------------------------------------------------ #

def should_win(g, p1, p2):
    """Does this pair of play symbols constitute a win, per the game's rule?
    Returns (should_win, forbidden_pattern)."""
    rule = g.cfg["win_rule"]
    sp = g.specials
    n_sp = (p1 in sp) + (p2 in sp)
    matched = (p1 == p2)
    forbidden = False
    f = g.cfg.get("forbidden")
    if f == "two_special" and n_sp == 2:
        forbidden = True
    if rule == "match_two":
        win = matched
    elif rule == "match_two_or_special":
        win = matched or n_sp == 1
    elif rule == "special_only":
        win = n_sp == 1
    else:
        win = matched
    return win, forbidden


# ------------------------------------------------------------------ #
#  Checks                                                             #
# ------------------------------------------------------------------ #

def check_01_play_positions(job, g):
    rule_text = {
        "match_two": "two similar symbols",
        "match_two_or_special": "two similar symbols or a single special symbol",
        "special_only": "a single special symbol (a matching pair does not win)",
    }[g.cfg["win_rule"]]
    r = CheckResult(1, "Play positions",
        "Check whether the fields pertaining to the play positions in the data "
        "layout have been populated as per the game specifications "
        f"(win = {rule_text}).")
    bad_symbol, bad_blank, bad_status, bad_forbidden = [], [], [], []
    for i in range(job.tickets):
        if not g.present(i):
            continue
        p1, p2 = g.p1[i], g.p2[i]
        if p1 not in g.symbols or p2 not in g.symbols:
            bad_symbol.append(g.serial[i]); continue
        if not g.blanks_ok[i]:
            bad_blank.append(g.serial[i]); continue
        win, forbidden = should_win(g, p1, p2)
        if forbidden:
            bad_forbidden.append(f"{g.serial[i]}: {p1}+{p2}")
            continue
        if g.is_winner(i) != win:
            bad_status.append(g.serial[i])
    if bad_symbol:
        r.fail(bad_symbol); r.note(f"{len(bad_symbol)} ticket(s) with an invalid play-symbol code.")
    if bad_blank:
        r.fail(bad_blank); r.note(f"{len(bad_blank)} ticket(s) where the unused play positions are not blank.")
    if bad_forbidden:
        r.fail(bad_forbidden)
        r.note(f"{len(bad_forbidden)} ticket(s) carrying a prohibited symbol combination.")
    if bad_status:
        r.fail(bad_status); r.note(f"{len(bad_status)} ticket(s) where winner/loser status does not match the play symbols.")
    if r.passed:
        r.note("No exceptions were noted. All play positions are valid and "
               "consistent with winner/loser status.")
    return r


def check_02_prize_box(job, g):
    r = CheckResult(2, "Prize box / position",
        "Check whether the field pertaining to the prize box/position in the "
        "data layout has been populated as per the game specifications.")
    if not g.cfg.get("has_prize_spot", True):
        bad = [g.serial[i] for i in range(job.tickets)
               if g.present(i) and g.pcode[i].strip() != ""]
        if bad:
            r.fail(bad); r.note(f"{len(bad)} ticket(s) carry a prize print code "
                                "although this game has no prize spot.")
        else:
            r.note("Not applicable: this game has no prize spot, and the prize "
                   "print code is blank on every ticket as specified.")
        return r
    valid = {p["print"] for p in g.cfg["prizes"] if p.get("print")}
    bad_sym, bad_map = [], []
    for i in range(job.tickets):
        if not g.present(i):
            continue
        pc = g.pcode[i]
        if pc not in valid:
            bad_sym.append(g.serial[i]); continue
        if g.is_winner(i):
            want = g.prize_by_sys[g.sys[i]]["print"]
            if pc != want:
                bad_map.append(g.serial[i])
    if bad_sym:
        r.fail(bad_sym); r.note(f"{len(bad_sym)} ticket(s) with an invalid prize print code.")
    if bad_map:
        r.fail(bad_map); r.note(f"{len(bad_map)} winning ticket(s) whose prize print code does not match the system prize code.")
    if r.passed:
        r.note("No exceptions were noted. All prize print codes are valid and, "
               "for winners, match the system prize code.")
    return r


def check_03_serial(job, g, info):
    r = CheckResult(3, "Serial number",
        "Check whether the serial-number field is populated per the format "
        "PXXXGGGGTTTTTTT with a unique consecutive number from 0000001 to "
        f"{job.tickets:07d}.")
    lot = g.cfg["lottery"]
    missing = [f"{job.flag}{lot}{g.key}{n+1:07d}"
               for n in range(job.tickets) if g.serial[n] is None]
    dups = [s for (k, s) in info["duplicates"] if k == g.key]
    chans = set(game_channels(g.cfg))
    mal = [f"row {rw}: {s}" for (rw, c, s) in info["malformed"] if c in chans]
    wrong = [f"row {rw}: {s}" for (rw, c, s) in info["wrong_game"] if c in chans]
    if missing:
        r.fail(missing); r.note(f"{len(missing):,} serial number(s) missing.")
    if dups:
        r.fail(dups); r.note(f"{len(dups)} duplicate serial(s).")
    if mal:
        r.fail(mal); r.note(f"{len(mal)} malformed serial(s).")
    if wrong:
        r.fail(wrong); r.note(f"{len(wrong)} serial(s) with the wrong game number.")
    if r.passed:
        r.note(f"No exceptions were noted. All {job.tickets:,} serials are "
               "present exactly once and correctly formatted.")
    return r


def check_04_virn_barcode(job, g):
    r = CheckResult(4, "VIRN and Barcode",
        "Check whether the VIRN and Barcode fields are populated per the format "
        "PXXXGGGGVVVVVVV, with the barcode carrying a 16th parity digit "
        "(parity = sum of the VIRN digits, mod 10).")
    lot = g.cfg["lottery"]
    seen = set()
    bad_fmt, bad_range, dup, bad_bar = [], [], [], []
    for i in range(job.tickets):
        if not g.present(i):
            continue
        v, b = g.virn[i], g.bar[i]
        if (len(v) != 15 or not v.isdigit() or v[0] != job.flag
                or v[1:4] != lot or v[4:8] != g.key):
            bad_fmt.append(g.serial[i]); continue
        rnd = int(v[8:])
        if not (job.virn_lo <= rnd <= job.virn_hi):
            bad_range.append(g.serial[i])
        if v in seen:
            dup.append(g.serial[i])
        else:
            seen.add(v)
        if b != v + str(sum(int(d) for d in v) % 10):
            bad_bar.append(g.serial[i])
    if bad_fmt:
        r.fail(bad_fmt); r.note(f"{len(bad_fmt)} VIRN(s) with an invalid format.")
    if bad_range:
        r.fail(bad_range); r.note(f"{len(bad_range)} VIRN(s) outside "
                                  f"{job.virn_lo:,}-{job.virn_hi:,}.")
    if dup:
        r.fail(dup); r.note(f"{len(dup)} duplicate VIRN(s).")
    if bad_bar:
        r.fail(bad_bar); r.note(f"{len(bad_bar)} barcode(s) with an incorrect parity digit.")
    if r.passed:
        r.note("No exceptions were noted. All VIRNs are valid, unique and their "
               "barcodes carry the correct parity digit.")
    return r


def check_05_agent_codes(job, g):
    r = CheckResult(5, "Agent validation codes",
        "Check whether the agent validation code fields are populated per the "
        "rules for high-tier/losing tickets and low-tier winning tickets.")
    bad_count, bad_letters = [], []
    for i in range(job.tickets):
        if not g.present(i):
            continue
        letters = [x for x in g.agent[i] if x]
        if len(letters) != job.n_agent:
            bad_count.append(g.serial[i]); continue
        low = job.low_letters.get(g.amount(i)) if g.is_winner(i) else None
        if low:
            others = [x for x in letters if x != low]
            ok = (letters.count(low) == 1 and len(others) == 2
                  and all(x in job.high_letters for x in others)
                  and others[0] != others[1])
        else:
            ok = (all(x in job.high_letters for x in letters)
                  and len(set(letters)) == len(letters))
        if not ok:
            bad_letters.append(f"{g.serial[i]}: {'/'.join(letters)}")
    if bad_count:
        r.fail(bad_count); r.note(f"{len(bad_count)} ticket(s) not carrying exactly {job.n_agent} agent letters.")
    if bad_letters:
        r.fail(bad_letters); r.note(f"{len(bad_letters)} ticket(s) whose agent letters break the composition rules.")
    if r.passed:
        r.note("No exceptions were noted. All agent validation codes follow the "
               "prescribed composition rules.")
    return r


def check_06_prize_structure(job, g):
    r = CheckResult(6, "Prize structure",
        "Check whether the prize structure prescribed by the game "
        "specification is reflected in the data layout.")
    counts = Counter(g.sys[i] for i in range(job.tickets) if g.present(i))
    ok_all = True
    pool = 0
    for p in sorted(g.cfg["prizes"], key=lambda p: p["amount"]):
        got = counts.get(p["sys"], 0)
        pool += got * p["amount"]
        ok = got == p["count"]
        ok_all &= ok
        label = p.get("label") or job.money(p["amount"])
        pr = p.get("print") or "-"
        r.note(f"{label} ({p['sys']}/{pr}, {p['tier']}): expected "
               f"{p['count']:,}, found {got:,}  {'OK' if ok else 'MISMATCH'}")
    loss = counts.get(g.loss_code, 0)
    ok = loss == g.cfg["loss"]["count"]; ok_all &= ok
    r.note(f"Losing tickets ({g.loss_code}): expected "
           f"{g.cfg['loss']['count']:,}, found {loss:,}  {'OK' if ok else 'MISMATCH'}")
    total = sum(counts.values())
    ok = total == job.tickets; ok_all &= ok
    r.note(f"Total tickets: expected {job.tickets:,}, found {total:,}  "
           f"{'OK' if ok else 'MISMATCH'}")
    winners = total - loss
    if winners:
        r.note(f"Winners {winners:,} (win ratio 1:{total/winners:.2f}); "
               f"prize pool {job.money(pool)}")
    if not ok_all:
        r.passed = False
        for p in g.cfg["prizes"]:
            if counts.get(p["sys"], 0) != p["count"]:
                r.offenders.append(f"{job.money(p['amount'])} ({p['sys']}): "
                                   f"found {counts.get(p['sys'],0)}, expected {p['count']}")
        if loss != g.cfg["loss"]["count"]:
            r.offenders.append(f"Loss: found {loss}, expected {g.cfg['loss']['count']}")
    return r


def check_07_gleps(job, g):
    need = g.cfg.get("gleps_per_book")
    amt = job.gleps_amount
    r = CheckResult(7, "GLEPS per book",
        f"Check whether each book of {job.book_size} tickets contains the "
        f"guaranteed {need} x {job.money(amt)} low-end prizes.")
    if not need:
        r.note("No GLEPS promise is configured for this game.")
        return r
    bad = []
    for bno, s, e in iter_groups(job, 1):
        got = count_amount(g, s, e, amt)
        if got != need:
            bad.append(f"Book {bno} (serials {s+1}-{e}): {got} (expected {need})")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} book(s) without exactly {need} x {job.money(amt)}.")
    else:
        r.note(f"No exceptions were noted. Every book contains exactly {need} x "
               f"{job.money(amt)} ({job.money(need*amt)} per book).")
    return r


def check_08_packing(job, g):
    sched = g.cfg.get("packing") or []
    parts = []
    for s in sched:
        a = ", ".join(f"{job.money(x['amount'])}"
                      f"{'>=' if x['mode']=='min' else '='}{x['n']}"
                      for x in s["allocation"])
        parts.append(f"{s['slots']} x {s['books']} books ({a})")
    r = CheckResult(8, "Packing allocation",
        "Check whether each packing slot carries the prescribed prize "
        "allocation: " + "; ".join(parts) + ".")
    if not sched:
        r.note("No packing schedule is configured for this game.")
        return r
    bad = []
    book_cursor = 0
    slot_no = 0
    for s in sched:
        for _ in range(s["slots"]):
            slot_no += 1
            start = book_cursor * job.book_size
            end = start + s["books"] * job.book_size
            book_cursor += s["books"]
            probs = []
            for a in s["allocation"]:
                got = count_amount(g, start, min(end, job.tickets), a["amount"])
                if a["mode"] == "exact" and got != a["n"]:
                    probs.append(f"{job.money(a['amount'])}={got} (expected {a['n']})")
                elif a["mode"] == "min" and got < a["n"]:
                    probs.append(f"{job.money(a['amount'])}={got} (min {a['n']})")
            if probs:
                bad.append(f"Slot {slot_no} (serials {start+1}-{min(end,job.tickets)}): "
                           + ", ".join(probs))
    if bad:
        r.fail(bad); r.note(f"{len(bad)} slot(s) not meeting the prescribed allocation.")
    else:
        r.note(f"No exceptions were noted. All {slot_no} packing slots meet the "
               "prescribed allocation.")
    return r


def check_09_group_rules(job, g):
    rules = g.cfg.get("group_rules") or []
    descs = []
    for gr in rules:
        lab = "/".join(job.money(a) for a in gr["amounts"])
        unit = "book" if gr.get("size", 1) == 1 else f"batch of {gr['size']} books"
        descs.append(f"each {unit}: {lab} {gr['mode'].replace('_',' ')} {gr.get('n',1)}")
    r = CheckResult(9, "Per-group prize promises",
        "Check the per-book and per-batch prize promises: "
        + "; ".join(descs) + "." if descs else
        "No per-group prize promises are configured for this game.")
    if not rules:
        r.note("Not applicable: this game declares no per-group prize promises.")
        return r
    bad = []
    for gr in rules:
        size = gr.get("size", 1)
        n = gr.get("n", 1)
        mode = gr["mode"]
        lab = "/".join(job.money(a) for a in gr["amounts"])
        unit = "Book" if size == 1 else f"Batch({size})"
        for idx, s, e in iter_groups(job, size):
            counts = {a: count_amount(g, s, e, a) for a in gr["amounts"]}
            tot = sum(counts.values())
            probs = []
            if mode == "exactly":
                for a in gr["amounts"]:
                    if counts[a] != n:
                        probs.append(f"{job.money(a)}={counts[a]} (expected {n})")
            elif mode == "min":
                if tot < n:
                    probs.append(f"{lab} total={tot} (min {n})")
            elif mode == "at_most_each":
                for a in gr["amounts"]:
                    if counts[a] > n:
                        probs.append(f"{job.money(a)}={counts[a]} (max {n})")
            elif mode == "at_least_one_total":
                if tot < 1:
                    probs.append(f"{lab}: none present")
                for a in gr["amounts"]:
                    if counts[a] > n:
                        probs.append(f"{job.money(a)}={counts[a]} (max {n})")
            elif mode == "exactly_one_total":
                if tot != 1:
                    probs.append(f"{lab} total={tot} (expected exactly 1)")
            if probs:
                bad.append(f"{unit} {idx} (serials {s+1}-{e}): " + ", ".join(probs))
                if len(bad) > 500:
                    break
        if len(bad) > 500:
            break
    if bad:
        r.fail(bad); r.note(f"{len(bad)} group(s) breaking a per-group prize promise.")
    else:
        r.note("No exceptions were noted. Every book and batch meets its "
               "prize promise.")
    return r


def check_10_nonglep(job, g):
    r = CheckResult(10, "Non-GLEPS per book",
        f"Check whether each book of {job.book_size} tickets contains at least "
        "one NON-GLEPS prize.")
    bad = []
    for bno, s, e in iter_groups(job, 1):
        if not any(g.is_winner(i) and g.amount(i) != job.gleps_amount
                   for i in range(s, e)):
            bad.append(f"Book {bno} (serials {s+1}-{e})")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} book(s) with no NON-GLEPS prize.")
    else:
        r.note("No exceptions were noted. Every book contains at least one "
               "NON-GLEPS prize.")
    return r


def check_11_consecutive_pattern(job, g):
    r = CheckResult(11, "Consecutive play pattern",
        "Check whether consecutive tickets contain the same play symbol pattern.")
    bad = []
    for i in range(job.tickets - 1):
        if not (g.present(i) and g.present(i + 1)):
            continue
        if (g.p1[i], g.p2[i]) == (g.p1[i+1], g.p2[i+1]):
            bad.append(f"{g.serial[i]} & {g.serial[i+1]}")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} instance(s), affecting {2*len(bad)} tickets.")
    else:
        r.note("No exceptions were noted. Consecutive tickets do not contain "
               "the same play symbol pattern in any instance.")
    return r


def check_12_consecutive_rotation(job, g):
    r = CheckResult(12, "Consecutive symbols (any rotation)",
        "Check whether consecutive tickets contain the same winning symbols in "
        "any rotation.")
    bad = []
    for i in range(job.tickets - 1):
        if not (g.present(i) and g.present(i + 1)):
            continue
        if {g.p1[i], g.p2[i]} == {g.p1[i+1], g.p2[i+1]}:
            bad.append(f"{g.serial[i]} & {g.serial[i+1]}")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} instance(s), affecting {2*len(bad)} tickets.")
    else:
        r.note("No exceptions were noted. Consecutive tickets do not contain "
               "the same winning symbols in any rotation in any instance.")
    return r


def check_13_block_winner(job, g):
    n = job.win_every
    r = CheckResult(13, f"Every {n} tickets contains a winner",
        f"Check whether each block of {n} tickets starting from ticket 001 in "
        "each book contains at least one prize.")
    bad = []
    for bno, s, e in iter_groups(job, 1):
        for off in range(0, job.book_size, n):
            bs, be = s + off, min(s + off + n, e)
            if not any(g.is_winner(i) for i in range(bs, be)):
                bad.append(f"Book {bno} tickets {off+1}-{off+n} "
                           f"(serials {bs+1}-{be})")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} block(s) of {n} with no winner.")
    else:
        r.note(f"No exceptions were noted. Every block of {n} tickets contains "
               "at least one prize.")
    return r


def check_14_value_run(job, g):
    mx = job.max_run
    r = CheckResult(14, "Prize value run length",
        f"Check whether more than {mx} consecutive winning tickets contain the "
        "same winning value / prize amount.")
    bad = []
    start = None
    run = 0
    prev = None
    for i in range(job.tickets):
        if not g.present(i) or not g.is_winner(i):
            if run > mx:
                bad.append(f"{job.money(g.amount(start))} x{run}: "
                           f"serials {g.serial[start]}..{g.serial[prev]}")
            start, run, prev = None, 0, None
            continue
        if run and g.sys[i] == g.sys[start]:
            run += 1
        else:
            if run > mx:
                bad.append(f"{job.money(g.amount(start))} x{run}: "
                           f"serials {g.serial[start]}..{g.serial[prev]}")
            start, run = i, 1
        prev = i
    if run > mx:
        bad.append(f"{job.money(g.amount(start))} x{run}: "
                   f"serials {g.serial[start]}..{g.serial[prev]}")
    if bad:
        r.fail(bad); r.note(f"{len(bad)} run(s) of more than {mx} identical prize values.")
    else:
        r.note(f"No exceptions were noted. No more than {mx} consecutive "
               "winning tickets share the same prize value.")
    return r


def check_15_virn_independence(job, g):
    r = CheckResult(15, "VIRN independence",
        "Check whether there is any logical relationship between the play "
        "patterns and the VIRN.")
    lo, hi = job.virn_lo, job.virn_hi
    span = hi - lo + 1
    win = [0] * 10
    tot = [0] * 10
    for i in range(job.tickets):
        if not g.present(i):
            continue
        v = g.virn[i]
        if len(v) != 15 or not v.isdigit():
            continue
        rnd = int(v[8:])
        if not (lo <= rnd <= hi):
            continue
        b = min((rnd - lo) * 10 // span, 9)
        tot[b] += 1
        if g.is_winner(i):
            win[b] += 1
    n_all, w_all = sum(tot), sum(win)
    if not n_all:
        r.note("No tickets available for the independence screen.")
        return r
    p = w_all / n_all
    chi2 = 0.0
    for b in range(10):
        ew, el = tot[b] * p, tot[b] * (1 - p)
        if ew > 0:
            chi2 += (win[b] - ew) ** 2 / ew
        if el > 0:
            chi2 += ((tot[b] - win[b]) - el) ** 2 / el
    r.note(f"Overall winner rate {p*100:.4f}%; chi-square across 10 VIRN bands "
           f"= {chi2:.2f} (critical value 27.88 at 0.1%).")
    if chi2 > 27.88:
        r.fail([f"chi-square {chi2:.2f} exceeds 27.88."])
        r.note("A relationship between the VIRN and the play outcome cannot be "
               "ruled out.")
    else:
        r.note("No exceptions were noted. The winner rate is independent of the VIRN.")
    return r


def run_all_checks(job, g, info):
    return [
        check_01_play_positions(job, g),
        check_02_prize_box(job, g),
        check_03_serial(job, g, info),
        check_04_virn_barcode(job, g),
        check_05_agent_codes(job, g),
        check_06_prize_structure(job, g),
        check_07_gleps(job, g),
        check_08_packing(job, g),
        check_09_group_rules(job, g),
        check_10_nonglep(job, g),
        check_11_consecutive_pattern(job, g),
        check_12_consecutive_rotation(job, g),
        check_13_block_winner(job, g),
        check_14_value_run(job, g),
        check_15_virn_independence(job, g),
    ]


# ------------------------------------------------------------------ #
#  Appendices                                                         #
# ------------------------------------------------------------------ #

def appendix_I(job, g):
    a, b = Counter(), Counter()
    for i in range(job.tickets):
        if g.present(i):
            a[g.p1[i]] += 1
            b[g.p2[i]] += 1
    return [(name, code, a.get(code, 0), b.get(code, 0),
             a.get(code, 0) + b.get(code, 0))
            for code, name in g.cfg["play_symbols"].items()]


def appendix_II(job, g):
    has = g.cfg.get("has_prize_spot", True)
    c = Counter()
    for i in range(job.tickets):
        if g.present(i):
            c[g.pcode[i] if has else g.sys[i]] += 1
    rows = []
    for p in sorted(g.cfg["prizes"], key=lambda p: p["amount"]):
        code = p["print"] if has else p["sys"]
        rows.append((p.get("label") or job.money(p["amount"]),
                     code or "-", c.get(code, 0)))
    if not has:
        rows.append(("Losing", g.loss_code, c.get(g.loss_code, 0)))
    return rows


def appendix_III(job, g):
    t = {}
    for i in range(job.tickets):
        if not g.present(i):
            continue
        if g.is_winner(i):
            cat = 0 if g.amount(i) == job.gleps_amount else 1
        else:
            cat = 2
        for x in g.agent[i]:
            if x:
                t.setdefault(x, [0, 0, 0])[cat] += 1
    return [(k, *t[k], sum(t[k])) for k in sorted(t)]


# ------------------------------------------------------------------ #
#  Reports                                                            #
# ------------------------------------------------------------------ #

def write_markdown(path, job, results, apps, info, layout_path):
    L = []
    W = L.append
    W(f"# {job.cfg['job']['title']} - Layout Audit Report")
    W("")
    W(f"Job         : `{job.cfg['job']['id']}`")
    W(f"Layout file : `{os.path.basename(layout_path)}`")
    W(f"Generated   : {datetime.datetime.now():%Y-%m-%d %H:%M:%S}")
    W("")
    W("## File load summary")
    W("")
    live = info["rows"] * job.channels - info["void_tickets"]
    W(f"- Data rows read : {info['rows']:,}")
    if info["pages"]:
        W(f"- Pages : {len(info['pages']):,}")
    W(f"- Live tickets : {live:,}")
    W(f"- VOID tickets ignored : {info['void_tickets']:,} "
      f"({info['void_rows_leading']:,} leading rows, "
      f"{info['void_rows_trailing']:,} trailing)")
    W(f"- Malformed serials : {len(info['malformed']):,}")
    W(f"- Duplicate serials : {len(info['duplicates']):,}")
    W(f"- Rows with an unexpected field count : {len(info['bad_field_count']):,}")
    if info["qa"]:
        W("- QA sample code values : "
          + ", ".join(f"{k}={v:,}" for k, v in sorted(info["qa"].items())))
    W("")
    W("Void padding is expected at both ends of the file and is not an "
      f"exception: VOID rows are identified by the serial value `{job.void}`.")
    W("")
    for key, res in results.items():
        g = job.games[key]
        n_exc = sum(1 for r in res if not r.passed)
        W("---")
        W(f"# Game {g['lottery']}-{key} - \"{g['name']}\"")
        W("")
        W(f"**{'ALL CHECKS PASSED' if not n_exc else str(n_exc)+' CHECK(S) WITH EXCEPTIONS'}**")
        W("")
        for r in res:
            status = ("PASS - No exceptions were noted" if r.passed
                      else f"EXCEPTIONS NOTED ({len(r.offenders)})")
            W(f"## {r.index}. {r.title} - {status}")
            W("")
            W(f"*Procedure:* {r.procedure}")
            W("")
            for ln in r.finding:
                W(f"- {ln}")
            W("")
        apI, apII, apIII = apps[key]
        W(f"## Appendix I - Play symbol distribution ({key})")
        W("")
        W("| Play symbol | Code | Position 1 | Position 2 | Total |")
        W("|---|---|---:|---:|---:|")
        t1 = t2 = tt = 0
        for name, code, a, b, tot in apI:
            W(f"| {name} | {code} | {a:,} | {b:,} | {tot:,} |")
            t1 += a; t2 += b; tt += tot
        W(f"| **Total** |  | **{t1:,}** | **{t2:,}** | **{tt:,}** |")
        W("")
        W(f"## Appendix II - Prize amount symbol distribution ({key})")
        W("")
        W("| Prize amount | Code | Total |")
        W("|---|---|---:|")
        for label, code, tot in apII:
            W(f"| {label} | {code} | {tot:,} |")
        W("")
        W(f"## Appendix III - Agent validation code distribution ({key})")
        W("")
        W("| Letter | GLEPS | Non-GLEPS | Losing | Total |")
        W("|---|---:|---:|---:|---:|")
        ca = cb = cc = ct = 0
        for letter, gl, ng, ls, tot in apIII:
            W(f"| {letter} | {gl:,} | {ng:,} | {ls:,} | {tot:,} |")
            ca += gl; cb += ng; cc += ls; ct += tot
        W(f"| **Total** | **{ca:,}** | **{cb:,}** | **{cc:,}** | **{ct:,}** |")
        W("")
    open(path, "w").write("\n".join(L))


def write_word(path, job, results, apps, info, layout_path):
    from docx import Document
    from docx.shared import RGBColor
    doc = Document()
    doc.add_heading(f"{job.cfg['job']['title']} - Layout Audit Report", level=0)
    p = doc.add_paragraph()
    p.add_run(f"Job: {job.cfg['job']['id']}\n").italic = True
    p.add_run(f"Layout file: {os.path.basename(layout_path)}\n").italic = True
    p.add_run(f"Generated: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}").italic = True
    live = info["rows"] * job.channels - info["void_tickets"]
    doc.add_paragraph(
        f"Data rows read: {info['rows']:,}   |   Live tickets: {live:,}   |   "
        f"VOID ignored: {info['void_tickets']:,}   |   "
        f"Malformed serials: {len(info['malformed']):,}   |   "
        f"Duplicate serials: {len(info['duplicates']):,}")
    for key, res in results.items():
        g = job.games[key]
        n_exc = sum(1 for r in res if not r.passed)
        doc.add_page_break()
        doc.add_heading(f'Game {g["lottery"]}-{key} - "{g["name"]}"', level=1)
        run = doc.add_paragraph().add_run(
            "ALL CHECKS PASSED" if not n_exc else f"{n_exc} check(s) with exceptions")
        run.bold = True
        run.font.color.rgb = (RGBColor(0, 0x80, 0) if not n_exc
                              else RGBColor(0xC0, 0, 0))
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for j, h in enumerate(["Index", "Procedures", "Findings"]):
            c = t.rows[0].cells[j]
            c.text = h
            for para in c.paragraphs:
                for rr in para.runs:
                    rr.bold = True
        for r in res:
            c = t.add_row().cells
            c[0].text = str(r.index)
            c[1].text = r.procedure
            c[2].text = ""
            rr = c[2].paragraphs[0].add_run(
                "No exceptions were noted." if r.passed
                else f"Exceptions noted ({len(r.offenders)}).")
            rr.bold = True
            rr.font.color.rgb = (RGBColor(0, 0x80, 0) if r.passed
                                 else RGBColor(0xC0, 0, 0))
            for ln in r.finding:
                c[2].add_paragraph(ln)
        apI, apII, apIII = apps[key]
        doc.add_heading("Appendix I - Play symbol distribution", level=2)
        t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"
        for j, h in enumerate(["Play symbol", "Code", "Position 1", "Position 2", "Total"]):
            t.rows[0].cells[j].text = h
        for name, code, a, b, tot in apI:
            c = t.add_row().cells
            c[0].text, c[1].text = name, code
            c[2].text, c[3].text, c[4].text = f"{a:,}", f"{b:,}", f"{tot:,}"
        doc.add_heading("Appendix II - Prize amount symbol distribution", level=2)
        t = doc.add_table(rows=1, cols=3); t.style = "Table Grid"
        for j, h in enumerate(["Prize amount", "Code", "Total"]):
            t.rows[0].cells[j].text = h
        for label, code, tot in apII:
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text = label, code, f"{tot:,}"
        doc.add_heading("Appendix III - Agent validation code distribution", level=2)
        t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"
        for j, h in enumerate(["Letter", "GLEPS", "Non-GLEPS", "Losing", "Total"]):
            t.rows[0].cells[j].text = h
        for letter, gl, ng, ls, tot in apIII:
            c = t.add_row().cells
            c[0].text = letter
            c[1].text, c[2].text, c[3].text, c[4].text = \
                f"{gl:,}", f"{ng:,}", f"{ls:,}", f"{tot:,}"
    doc.save(path)


def write_csvs(out_dir, results):
    import csv
    written = []
    for key, res in results.items():
        for r in res:
            if r.passed or not r.offenders:
                continue
            fp = os.path.join(out_dir, f"game{key}_check{r.index:02d}_offenders.csv")
            with open(fp, "w", newline="") as f:
                w = csv.writer(f)
                w.writerow(["game", "check", "title", "item"])
                for o in r.offenders:
                    w.writerow([key, r.index, r.title, o])
            written.append(fp)
    return written


def write_summary_json(path, job, results, info, layout_path, files):
    """Machine-readable summary for callers that do not parse the reports."""
    live = info["rows"] * job.channels - info["void_tickets"]
    games = []
    for key, res in results.items():
        games.append({
            "key": key,
            "name": job.games[key]["name"],
            "passed": all(r.passed for r in res),
            "checks": [{"index": r.index, "title": r.title, "passed": r.passed,
                        "offenders": len(r.offenders)} for r in res],
        })
    summary = {
        "job_id": job.cfg["job"]["id"],
        "layout_file": os.path.basename(layout_path),
        "generated": f"{datetime.datetime.now():%Y-%m-%d %H:%M:%S}",
        "load": {
            "rows": info["rows"],
            "live_tickets": live,
            "void_tickets": info["void_tickets"],
            "malformed_serials": len(info["malformed"]),
            "duplicate_serials": len(info["duplicates"]),
            "bad_field_count_rows": len(info["bad_field_count"]),
        },
        "all_passed": all(g["passed"] for g in games),
        "exceptions": sum(1 for g in games for c in g["checks"] if not c["passed"]),
        "games": games,
        "files": [os.path.basename(f) for f in files],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=1)


# ------------------------------------------------------------------ #
#  Main                                                               #
# ------------------------------------------------------------------ #

def main():
    ap = argparse.ArgumentParser(description="NLB layout audit engine")
    ap.add_argument("--config", required=True)
    ap.add_argument("--layout", required=True)
    ap.add_argument("--out", default=".")
    ap.add_argument("--skip-validation", action="store_true")
    args = ap.parse_args()

    cfg = json.load(open(args.config))
    if not args.skip_validation:
        rep = validate(cfg)
        if not rep.ok:
            print("Settings failed validation - audit not run:")
            for game, msg in rep.errors:
                print(f"  X [{game}] {msg}")
            sys.exit(1)
        print(f"Settings validated: {len(cfg['games'])} game(s) consistent.")

    job = Job(cfg)
    os.makedirs(args.out, exist_ok=True)

    print(f"Loading layout: {args.layout}")
    games, info = load_layout(job, args.layout)
    live = info["rows"] * job.channels - info["void_tickets"]
    print(f"  data rows        : {info['rows']:,}")
    print(f"  live tickets     : {live:,}")
    print(f"  void tickets     : {info['void_tickets']:,} "
          f"({info['void_rows_leading']:,} leading rows, "
          f"{info['void_rows_trailing']:,} trailing)")
    print(f"  malformed serials: {len(info['malformed']):,}")
    print(f"  duplicate serials: {len(info['duplicates']):,}")

    results, apps = {}, {}
    for key in job.games:
        print(f"Auditing {key} ({job.games[key]['name']}) ...")
        g = games[key]
        results[key] = run_all_checks(job, g, info)
        apps[key] = (appendix_I(job, g), appendix_II(job, g), appendix_III(job, g))

    md = os.path.join(args.out, "audit_report.md")
    dx = os.path.join(args.out, "audit_report.docx")
    write_markdown(md, job, results, apps, info, args.layout)
    try:
        write_word(dx, job, results, apps, info, args.layout)
    except ImportError:
        dx = None
        print("  (python-docx not installed - Word report skipped)")
    csvs = write_csvs(args.out, results)
    sj = os.path.join(args.out, "audit_summary.json")
    write_summary_json(sj, job, results, info, args.layout,
                       [md] + ([dx] if dx else []) + csvs)

    print("\nReports written:")
    print(f"  {md}")
    if dx:
        print(f"  {dx}")
    for c in csvs:
        print(f"  {c}")
    print(f"  {sj}")

    print("\n================ SUMMARY ================")
    for key, res in results.items():
        exc = [r for r in res if not r.passed]
        print(f"Game {key} ({job.games[key]['name']}): "
              f"{'ALL PASSED' if not exc else str(len(exc))+' exception(s)'}")
        for r in exc:
            print(f"    - Check {r.index} {r.title}: {len(r.offenders)} item(s)")


if __name__ == "__main__":
    main()
